import os

from docx import Document
from pptx import Presentation
from reportlab.lib.pagesizes import A4
from reportlab.pdfgen import canvas


# ---------------------------------------
# Output folder
# ---------------------------------------

OUTPUT_DIR = "generated_outputs"

os.makedirs(OUTPUT_DIR, exist_ok=True)


# ---------------------------------------
# Person 5 - Multi Format Output Engine
# ---------------------------------------

class MultiFormatOutputEngine:

    def __init__(self, content):
        self.content = content

    # -----------------------------------
    # TXT Generation
    # -----------------------------------

    def create_txt(self):

        file_path = os.path.join(
            OUTPUT_DIR,
            "output.txt"
        )

        with open(
            file_path,
            "w",
            encoding="utf-8"
        ) as file:

            file.write(
                self.content["title"] + "\n"
            )

            file.write("=" * 50 + "\n\n")

            for section in self.content["sections"]:

                file.write(
                    section["heading"] + "\n"
                )

                file.write("-" * 30 + "\n")

                for point in section["points"]:

                    file.write(
                        "• " + point + "\n"
                    )

                file.write("\n")

        return file_path

    # -----------------------------------
    # DOCX Generation
    # -----------------------------------

    def create_docx(self):

        file_path = os.path.join(
            OUTPUT_DIR,
            "output.docx"
        )

        document = Document()

        document.add_heading(
            self.content["title"],
            level=0
        )

        for section in self.content["sections"]:

            document.add_heading(
                section["heading"],
                level=1
            )

            for point in section["points"]:

                document.add_paragraph(
                    point,
                    style="List Bullet"
                )

        document.save(file_path)

        return file_path

    # -----------------------------------
    # PPTX Generation
    # -----------------------------------

    def create_ppt(self):

        file_path = os.path.join(
            OUTPUT_DIR,
            "output.pptx"
        )

        presentation = Presentation()

        # Title slide
        slide = presentation.slides.add_slide(
            presentation.slide_layouts[0]
        )

        slide.shapes.title.text = (
            self.content["title"]
        )

        slide.placeholders[1].text = (
            "AI Content Transformation Platform"
        )

        # Content slides
        for section in self.content["sections"]:

            slide = presentation.slides.add_slide(
                presentation.slide_layouts[1]
            )

            slide.shapes.title.text = (
                section["heading"]
            )

            text_frame = (
                slide.placeholders[1].text_frame
            )

            text_frame.clear()

            for index, point in enumerate(
                section["points"]
            ):

                if index == 0:
                    paragraph = text_frame.paragraphs[0]
                else:
                    paragraph = text_frame.add_paragraph()

                paragraph.text = point
                paragraph.level = 0

        presentation.save(file_path)

        return file_path

    # -----------------------------------
    # PDF Generation
    # -----------------------------------

    def create_pdf(self):

        file_path = os.path.join(
            OUTPUT_DIR,
            "output.pdf"
        )

        pdf = canvas.Canvas(
            file_path,
            pagesize=A4
        )

        width, height = A4

        y = height - 50

        # Title
        pdf.setFont(
            "Helvetica-Bold",
            20
        )

        pdf.drawString(
            50,
            y,
            self.content["title"]
        )

        y -= 45

        # Sections
        for section in self.content["sections"]:

            # Check page space
            if y < 80:

                pdf.showPage()

                y = height - 50

            pdf.setFont(
                "Helvetica-Bold",
                14
            )

            pdf.drawString(
                50,
                y,
                section["heading"]
            )

            y -= 25

            pdf.setFont(
                "Helvetica",
                10
            )

            for point in section["points"]:

                # Simple text wrapping
                words = point.split()

                line = ""

                for word in words:

                    test_line = (
                        line + " " + word
                    ).strip()

                    if len(test_line) > 85:

                        pdf.drawString(
                            65,
                            y,
                            "• " + line
                        )

                        y -= 18

                        line = word

                    else:

                        line = test_line

                if line:

                    pdf.drawString(
                        65,
                        y,
                        "• " + line
                    )

                    y -= 18

                # New page
                if y < 50:

                    pdf.showPage()

                    y = height - 50

                    pdf.setFont(
                        "Helvetica",
                        10
                    )

            y -= 15

        pdf.save()

        return file_path

    # -----------------------------------
    # Generate all formats
    # -----------------------------------

    def generate_all(self):

        files = {}

        files["TXT"] = self.create_txt()

        files["DOCX"] = self.create_docx()

        files["PPTX"] = self.create_ppt()

        files["PDF"] = self.create_pdf()

        return files


# =======================================
# MAIN PROGRAM
# =======================================

if __name__ == "__main__":

    # Sample validated content
    # In the real project, this data
    # comes from the previous team members.

    validated_content = {

        "title":
            "Artificial Intelligence in Healthcare",

        "sections": [

            {
                "heading":
                    "Introduction",

                "points": [

                    "Artificial Intelligence is increasingly used in healthcare.",

                    "AI can assist doctors in analysing medical information.",

                    "AI technology can support healthcare professionals."
                ]
            },

            {
                "heading":
                    "Applications",

                "points": [

                    "Medical image analysis",

                    "Patient data analysis",

                    "Clinical decision support",

                    "Healthcare workflow automation"
                ]
            },

            {
                "heading":
                    "Benefits",

                "points": [

                    "Faster analysis of information",

                    "Improved decision support",

                    "Efficient healthcare workflows",

                    "Better access to useful information"
                ]
            },

            {
                "heading":
                    "Conclusion",

                "points": [

                    "Artificial Intelligence can support modern healthcare.",

                    "Human professionals remain important for final decisions."
                ]
            }
        ]
    }

    # Create engine
    engine = MultiFormatOutputEngine(
        validated_content
    )

    # Generate files
    generated_files = engine.generate_all()

    # Display result
    print()
    print("=" * 55)
    print("       PERSON 5 - OUTPUT ENGINE")
    print("=" * 55)

    print()
    print("Output generated successfully!")
    print()

    for file_type, file_path in generated_files.items():

        print(
            f"{file_type:<6} -> {file_path}"
        )

    print()
    print("=" * 55)
    print("All files are available inside:")
    print(f"    {OUTPUT_DIR}/")
    print("=" * 55)